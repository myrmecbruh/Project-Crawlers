"""The master spreadsheet: how it is laid out, read, written and reconciled.

docs/crawlers.xlsx is the AUTHORITY on every number and name in the game.
src/defaults.json carries the same values so a fresh checkout still builds.

The one failure this arrangement exists to prevent is a value that lives in only
one of the two places: a knob in the sheet the game never reads, or a number in
the code the sheet cannot reach. reconcile() treats that as a build error, not a
warning, because a dial that is not connected to anything is worse than no dial.
"""

import collections
import json
import pathlib

# section -> (id column heading, value column headings, read-only?)
# The first value column is what the game reads for the simple sections.
SECTIONS = {
    "knobs":    ("key", ["value", "unit", "note"], False),
    "geometry": ("key", ["value", "unit", "note"], True),
    "names":    ("key", ["text"], False),
    "tags":     ("id",  ["name", "note"], False),
    "tiles":    ("id",  ["name", "tags", "top", "side", "footing", "cross", "clear", "note"], False),
    "attributes": ("id", ["name", "abbrev", "note"], False),
    "skills":   ("id",  ["name", "derives", "needs_tag", "without", "note"], False),
    "bones":    ("id",  ["name", "parent", "x", "y", "z", "note"], False),
    "slots":    ("id",  ["name", "note"], False),
    "speeds":   ("id",  ["name", "multiplier", "note"], False),
    "gear":     ("id",  ["name", "slot", "tags", "attr", "bonus", "note"], False),
    "figure":   ("id",  ["name", "slot", "item", "bone", "colour", "ox", "oy",
                         "from_m", "to_m", "w_top", "w_bot", "d_top", "d_bot",
                         "sides", "rings", "bulge", "cap_top", "cap_bot",
                         "note"], False),
    "structures": ("id", ["name", "tags", "colour", "half_width", "height_m",
                          "skill", "difficulty", "note"], False),
}

# Columns that are documentation. They may be edited freely and are never
# reported as a moved value.
PROSE = {"note", "unit"}


def load_defaults(path):
    raw = json.loads(pathlib.Path(path).read_text())
    return {s: raw[s] for s in SECTIONS}


def blank(v):
    return v is None or (isinstance(v, str) and not v.strip())


def norm(col, v):
    """Compare 480 and 480.0 as the same number; compare text with the ends
    trimmed, because a spreadsheet loves a trailing space."""
    if blank(v):
        return ""
    if col in ("value", "cross", "clear", "half_width", "from_m", "to_m",
               "height_m", "difficulty", "without", "multiplier",
               "x", "y", "z", "ox", "oy",
               "w_top", "w_bot", "d_top", "d_bot",
               "sides", "rings", "bulge", "cap_top", "cap_bot"):
        try:
            f = float(v)
            return int(f) if f == int(f) else round(f, 10)
        except (TypeError, ValueError):
            return str(v).strip()
    if col in ("footing", "slot", "item", "bone", "parent", "needs_tag"):
        return str(v).strip().lower()
    if col in ("derives", "attr", "bonus"):
        return ",".join(t.strip().lower() for t in str(v).split(",") if t.strip())
    if col == "tags":
        return ",".join(t.strip() for t in str(v).split(",") if t.strip())
    return str(v).strip()


def reconcile(defaults, sheet):
    """Merge the sheet over the defaults.

    Returns (merged, moved, errors). `moved` is the receipt printed on every
    build: every value the sheet actually changed. `errors` is fatal.
    """
    merged, moved, errors = {}, [], []

    for section, (idcol, cols, readonly) in SECTIONS.items():
        d = defaults.get(section, {})
        s = sheet.get(section, {}) if sheet else {}
        merged[section] = {}

        if sheet is not None:
            only_sheet = sorted(set(s) - set(d))
            only_code = sorted(set(d) - set(s))
            for k in only_sheet:
                errors.append("%s: '%s' is in the spreadsheet but not in the game "
                              "(a dial connected to nothing)" % (section, k))
            for k in only_code:
                errors.append("%s: '%s' is in the game but not in the spreadsheet "
                              "(a number nobody can reach)" % (section, k))

        for key, drow in d.items():
            row = dict(drow)
            srow = s.get(key) if sheet else None
            if srow:
                for col in cols:
                    if blank(srow.get(col)):
                        continue          # an empty cell means "leave it alone"
                    a, b = norm(col, drow.get(col)), norm(col, srow.get(col))
                    row[col] = srow[col]
                    if a != b and col not in PROSE:
                        if readonly:
                            errors.append(
                                "geometry: '%s' was changed in the spreadsheet "
                                "(%s -> %s). Geometry is recorded, never tuned."
                                % (key, a, b))
                        else:
                            moved.append((section, key, col, a, b))
            merged[section][key] = row

    return merged, moved, errors


def for_game(merged):
    """Strip the spreadsheet down to what the game actually reads."""
    out = {
        "knobs":    {k: norm("value", r["value"]) for k, r in merged["knobs"].items()},
        "geometry": {k: norm("value", r["value"]) for k, r in merged["geometry"].items()},
        "names":    {k: str(r["text"]) for k, r in merged["names"].items()},
        "tags":     {k: {"name": str(r["name"]), "note": str(r.get("note", ""))}
                     for k, r in merged["tags"].items()},
        "tiles":    {},
    }
    for k, r in merged["tiles"].items():
        out["tiles"][k] = {
            "name": str(r["name"]),
            "tags": [t for t in norm("tags", r["tags"]).split(",") if t],
            "top": str(r["top"]).strip(),
            "side": str(r["side"]).strip(),
            "footing": norm("footing", r["footing"]),
            "cross": norm("cross", r["cross"]),
            "clear": norm("clear", r["clear"]),
            "note": str(r.get("note", "")),
        }

    out["attributes"] = {k: {"name": str(r["name"]), "abbrev": str(r["abbrev"]).strip(),
                             "note": str(r.get("note", ""))}
                         for k, r in merged["attributes"].items()}

    # "agility:2,endurance:1" -> [["agility", 2.0], ["endurance", 1.0]]
    out["skills"] = {}
    for k, r in merged["skills"].items():
        pairs = []
        for part in norm("derives", r["derives"]).split(","):
            if not part:
                continue
            attr, _, w = part.partition(":")
            pairs.append([attr.strip(), float(w) if w.strip() else 1.0])
        out["skills"][k] = {"name": str(r["name"]), "derives": pairs,
                            "needs_tag": norm("needs_tag", r.get("needs_tag")),
                            "without": norm("without", r.get("without")) or 0,
                            "note": str(r.get("note", ""))}

    out["structures"] = {k: {"name": str(r["name"]),
                             "tags": [t for t in norm("tags", r["tags"]).split(",") if t],
                             "colour": str(r["colour"]).strip(),
                             "half_width": norm("half_width", r["half_width"]),
                             "height_m": norm("height_m", r["height_m"]),
                             "skill": str(r["skill"]).strip().lower(),
                             "difficulty": norm("difficulty", r["difficulty"]),
                             "note": str(r.get("note", ""))}
                         for k, r in merged["structures"].items()}

    out["bones"] = {k: {"name": str(r["name"]), "parent": norm("parent", r["parent"]),
                        "x": norm("x", r["x"]), "y": norm("y", r["y"]),
                        "z": norm("z", r["z"]), "note": str(r.get("note", ""))}
                    for k, r in merged["bones"].items()}
    out["slots"] = {k: {"name": str(r["name"]), "note": str(r.get("note", ""))}
                    for k, r in merged["slots"].items()}
    out["speeds"] = {k: {"name": str(r["name"]),
                         "multiplier": norm("multiplier", r["multiplier"]),
                         "note": str(r.get("note", ""))}
                     for k, r in merged["speeds"].items()}
    def pairs_of(text):
        out_ = []
        if blank(text):
            return out_
        for part_ in norm("attr", text).split(","):
            if not part_:
                continue
            key, _, amt = part_.partition(":")
            try:
                out_.append([key.strip(), float(amt)])
            except ValueError:
                out_.append([key.strip(), 0.0])
        return out_

    out["gear"] = {k: {"name": str(r["name"]), "slot": norm("slot", r["slot"]),
                       "tags": [t for t in norm("tags", r["tags"]).split(",") if t],
                       "attr": pairs_of(r.get("attr")), "bonus": pairs_of(r.get("bonus")),
                       "note": str(r.get("note", ""))}
                   for k, r in merged["gear"].items()}
    out["figure"] = {k: {"name": str(r["name"]), "slot": norm("slot", r["slot"]),
                         "item": norm("item", r["item"]),
                         "bone": norm("bone", r["bone"]),
                         "colour": str(r["colour"]).strip(),
                         "ox": norm("ox", r["ox"]), "oy": norm("oy", r["oy"]),
                         "from_m": norm("from_m", r["from_m"]),
                         "to_m": norm("to_m", r["to_m"]),
                         "w_top": norm("w_top", r["w_top"]),
                         "w_bot": norm("w_bot", r["w_bot"]),
                         "d_top": norm("d_top", r["d_top"]),
                         "d_bot": norm("d_bot", r["d_bot"]),
                         "sides": int(norm("sides", r.get("sides")) or 6),
                         "rings": int(norm("rings", r.get("rings")) or 2),
                         "bulge": norm("bulge", r.get("bulge")) or 0,
                         "cap_top": norm("cap_top", r.get("cap_top")) or 0,
                         "cap_bot": norm("cap_bot", r.get("cap_bot")) or 0,
                         "note": str(r.get("note", ""))}
                     for k, r in merged["figure"].items()}
    return out


def check_vocabulary(game):
    """Every tag a tile claims must exist in the tag list, and every footing must
    be one the renderer knows. A typo in the sheet should stop the build, not
    quietly produce a tile with a tag nothing recognises."""
    errors = []
    known = set(game["tags"])
    for tid, t in game["tiles"].items():
        for tag in t["tags"]:
            if tag not in known:
                errors.append("tiles: '%s' claims the tag '%s', which is not in "
                              "the tags sheet" % (tid, tag))
        if t["footing"] not in ("walk", "ramp", "block"):
            errors.append("tiles: '%s' has footing '%s'; expected walk, ramp or "
                          "block" % (tid, t["footing"]))
        if not isinstance(t["cross"], (int, float)):
            errors.append("tiles: '%s' has a non-numeric crossing difficulty "
                          "'%s'" % (tid, t["cross"]))

    # Rule 2: every skill comes from the natural attributes. A skill that names
    # an attribute nobody has is a second framework starting to grow.
    attrs = set(game["attributes"])
    if len(attrs) != 6:
        errors.append("attributes: there are %d, and rule 2 says six" % len(attrs))
    for sid, sk in game["skills"].items():
        if not sk["derives"]:
            errors.append("skills: '%s' derives from no attribute at all "
                          "(rule 2)" % sid)
        for attr, w in sk["derives"]:
            if attr not in attrs:
                errors.append("skills: '%s' derives from '%s', which is not one "
                              "of the six attributes" % (sid, attr))
            if w <= 0:
                errors.append("skills: '%s' weights '%s' at %s" % (sid, attr, w))

    # Rule from the user: a skill is one PAIR of the six attributes, and every
    # pair exists exactly once. That makes the skill list a complete, checkable
    # grid rather than a list somebody keeps adding to.
    seen = {}
    for sid, sk in game["skills"].items():
        if len(sk["derives"]) != 2:
            errors.append("skills: '%s' is built from %d attributes; every skill "
                          "is a pair" % (sid, len(sk["derives"])))
            continue
        pair = tuple(sorted(a for a, _w in sk["derives"]))
        if pair[0] == pair[1]:
            errors.append("skills: '%s' pairs '%s' with itself" % (sid, pair[0]))
        elif pair in seen:
            errors.append("skills: '%s' and '%s' are both %s + %s"
                          % (seen[pair], sid, pair[0], pair[1]))
        else:
            seen[pair] = sid
    if attrs and len(attrs) == 6:
        want = set()
        ordered = list(game["attributes"])
        for i in range(len(ordered)):
            for j in range(i + 1, len(ordered)):
                want.add(tuple(sorted((ordered[i], ordered[j]))))
        for pair in sorted(want - set(seen)):
            errors.append("skills: nothing is %s + %s yet; every pair of the six "
                          "is a skill" % pair)

    known_skills = set(game["skills"])
    for stid, st in game["structures"].items():
        if st["skill"] not in known_skills:
            errors.append("structures: '%s' is raised with '%s', which is not a "
                          "skill" % (stid, st["skill"]))
        for tag in st["tags"]:
            if tag not in known:
                errors.append("structures: '%s' claims the tag '%s', which is not "
                              "in the tags sheet" % (stid, tag))

    # Low numbers is a rule, so the build keeps an eye on the scale.
    for kname, kmax in (("actor.attribute_max", 12), ("skill.cap", 12),
                        ("roll.die_faces", 12), ("roll.dice", 5)):
        v = game["knobs"].get(kname)
        if v is not None and v > kmax:
            errors.append("knobs: '%s' is %s. Rule 10 says low numbers; above "
                          "%s is not a number anybody holds in their head"
                          % (kname, v, kmax))
    dice = game["knobs"].get("roll.dice")
    if dice is not None and dice < 1:
        errors.append("knobs: 'roll.dice' is %s; an attempt needs at least one "
                      "die" % dice)

    speeds = list(game["speeds"].items())
    if len(speeds) < 2:
        errors.append("speeds: there is only one speed, so the control would do "
                      "nothing")
    for sid, sp in speeds:
        if not (sp["multiplier"] > 0):
            errors.append("speeds: '%s' runs at %s, which is not a speed"
                          % (sid, sp["multiplier"]))
    if speeds and speeds[0][1]["multiplier"] != 1:
        errors.append("speeds: the first row is '%s' at %sx; the first row is "
                      "normal speed and must be 1"
                      % (speeds[0][0], speeds[0][1]["multiplier"]))
    seen_mult = set()
    for sid, sp in speeds:
        if sp["multiplier"] in seen_mult:
            errors.append("speeds: two rows both run at %sx" % sp["multiplier"])
        seen_mult.add(sp["multiplier"])

    # A crawler is twelve parts, which are their gear. That is a rule, so the
    # build checks it rather than trusting the sheet to stay that way.
    if len(game["slots"]) != 12:
        errors.append("slots: there are %d gear slots, and a crawler has twelve"
                      % len(game["slots"]))

    bones = set(game["bones"])
    for bid, b in game["bones"].items():
        if b["parent"] and b["parent"] not in bones:
            errors.append("bones: '%s' hangs off '%s', which is not a bone"
                          % (bid, b["parent"]))
    # A loop in a skeleton hangs the game rather than failing it.
    for bid in game["bones"]:
        seen, cur = set(), bid
        while cur:
            if cur in seen:
                errors.append("bones: '%s' is part of a loop" % bid)
                break
            seen.add(cur)
            cur = game["bones"].get(cur, {}).get("parent", "")

    for sid, sk in game["skills"].items():
        if sk["needs_tag"] and sk["needs_tag"] not in known:
            errors.append("skills: '%s' needs gear tagged '%s', which is not in "
                          "the tags sheet" % (sid, sk["needs_tag"]))
        if sk["without"] > 0:
            errors.append("skills: '%s' has a penalty of +%s for going without, "
                          "which would be a reward" % (sid, sk["without"]))
        if sk["without"] and not sk["needs_tag"]:
            errors.append("skills: '%s' penalises going without, but never says "
                          "without what" % sid)

    slots = set(game["slots"])
    for gid, g in game["gear"].items():
        for attr, amt in g["attr"]:
            if attr not in attrs:
                errors.append("gear: '%s' shifts '%s', which is not one of the six"
                              % (gid, attr))
        for skill, amt in g["bonus"]:
            if skill not in game["skills"]:
                errors.append("gear: '%s' helps with '%s', which is not a skill"
                              % (gid, skill))
        if g["slot"] not in slots:
            errors.append("gear: '%s' goes in slot '%s', which is not one of the "
                          "twelve" % (gid, g["slot"]))
        for tag in g["tags"]:
            if tag not in known:
                errors.append("gear: '%s' claims the tag '%s', which is not in "
                              "the tags sheet" % (gid, tag))

    drawn = set()
    for fid, f in game["figure"].items():
        if f["slot"] != "body" and f["slot"] not in slots:
            errors.append("figure: '%s' has slot '%s'; expected body or one of "
                          "the twelve" % (fid, f["slot"]))
        if f["bone"] not in bones:
            errors.append("figure: '%s' hangs off '%s', which is not a bone"
                          % (fid, f["bone"]))
        if not (f["to_m"] > f["from_m"]):
            errors.append("figure: '%s' spans %s to %s, which is nothing at all"
                          % (fid, f["from_m"], f["to_m"]))
        for c in ("w_top", "w_bot", "d_top", "d_bot"):
            if not (f[c] > 0):
                errors.append("figure: '%s' has %s of %s" % (fid, c, f[c]))
        if f["sides"] < 3 or f["sides"] > 16:
            errors.append("figure: '%s' has %d sides; 3 is the fewest that is a "
                          "shape and 16 is as many as is worth drawing"
                          % (fid, f["sides"]))
        if f["rings"] < 2 or f["rings"] > 12:
            errors.append("figure: '%s' has %d rings along it; 2 is a straight "
                          "taper and 12 is plenty" % (fid, f["rings"]))
        # With only two rings both of them sit inside the rounded ends, so both
        # collapse to a point and the part disappears entirely. It takes a third
        # ring in the middle to have anything left.
        if (f["cap_top"] > 0 or f["cap_bot"] > 0) and f["rings"] < 3:
            errors.append("figure: '%s' rounds its ends off but has only %d "
                          "rings; a rounded end needs at least 3 or the whole "
                          "part collapses to a line" % (fid, f["rings"]))
        if f["cap_top"] + f["cap_bot"] > 1:
            errors.append("figure: '%s' rounds off %s of its length at the ends, "
                          "which is more than it has"
                          % (fid, f["cap_top"] + f["cap_bot"]))
        if f["bulge"] < -0.9:
            errors.append("figure: '%s' bulges %s, which turns it inside out"
                          % (fid, f["bulge"]))
        if f["slot"] == "body":
            if f["item"]:
                errors.append("figure: '%s' is body and should name no item" % fid)
        elif f["item"] not in game["gear"]:
            errors.append("figure: '%s' draws the gear '%s', which does not exist"
                          % (fid, f["item"]))
        elif game["gear"][f["item"]]["slot"] != f["slot"]:
            errors.append("figure: '%s' is in slot '%s' but draws '%s', which is "
                          "worn on '%s'" % (fid, f["slot"], f["item"],
                                            game["gear"][f["item"]]["slot"]))
        else:
            drawn.add(f["item"])

    # Rule 4: if it is on a crawler it is on the screen. Gear nothing draws
    # would be invisible when worn, which is exactly the failure rule 4 forbids.
    for gid in game["gear"]:
        if gid not in drawn:
            errors.append("gear: nothing draws '%s', so wearing it would show "
                          "nothing (rule 4)" % gid)
    for sid in slots:
        if not any(g["slot"] == sid for g in game["gear"].values()):
            errors.append("slots: nothing can be worn in '%s' yet" % sid)
    return errors


def read_xlsx(path):
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True)
    out = {}
    for section, (idcol, cols, _ro) in SECTIONS.items():
        if section not in wb.sheetnames:
            continue
        ws = wb[section]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        head = [str(c).strip().lower() if c is not None else "" for c in rows[0]]
        if idcol not in head:
            continue
        idx = {c: head.index(c) for c in [idcol] + cols if c in head}
        table = {}
        for r in rows[1:]:
            key = r[idx[idcol]] if idx[idcol] < len(r) else None
            if blank(key):
                continue
            row = {c: (r[i] if i < len(r) else None)
                   for c, i in idx.items() if c != idcol}
            # A row with nothing in any value column is a note to the reader
            # (the READ ONLY banner, a spacer), not a setting.
            if all(blank(v) for v in row.values()):
                continue
            table[str(key).strip()] = row
        out[section] = table
    return out


def write_xlsx(defaults, path, readme_lines):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "read me first"
    ws.column_dimensions["A"].width = 100
    for i, line in enumerate(readme_lines, start=1):
        ws.cell(row=i, column=1, value=line).alignment = Alignment(wrap_text=True, vertical="top")
        if line and not line.startswith(" ") and line == line.upper():
            ws.cell(row=i, column=1).font = Font(bold=True)

    head_fill = PatternFill("solid", fgColor="1F2A33")
    ro_fill = PatternFill("solid", fgColor="5C3A1E")

    for section, (idcol, cols, readonly) in SECTIONS.items():
        ws = wb.create_sheet(section)
        headings = [idcol] + cols
        for c, h in enumerate(headings, start=1):
            cell = ws.cell(row=1, column=c, value=h)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = ro_fill if readonly else head_fill
        for r, (key, row) in enumerate(defaults[section].items(), start=2):
            ws.cell(row=r, column=1, value=key)
            for c, col in enumerate(cols, start=2):
                ws.cell(row=r, column=c, value=row.get(col))
        widths = {"key": 26, "id": 16, "value": 10, "unit": 10, "note": 78,
                  "text": 54, "name": 18, "tags": 30, "top": 10, "side": 10,
                  "footing": 10}
        for c, h in enumerate(headings, start=1):
            ws.column_dimensions[get_column_letter(c)].width = widths.get(h, 16)
        ws.freeze_panes = "B2"
        if readonly:
            ws.cell(row=len(defaults[section]) + 3, column=1,
                    value="READ ONLY -- these describe the shape of the world, "
                          "not its balance. Changing one stops the build on purpose.")

    wb.save(path)


def selftest():
    """A broken ruler is worse than no ruler: prove reconcile() can actually
    catch the failures it exists to catch, before trusting it with a build."""
    checks = []

    def add(name, ok, detail=""):
        checks.append((name, ok, detail))

    base = {
        "knobs": {"a": {"value": 1, "unit": "x", "note": "n"}},
        "geometry": {"g": {"value": 2, "unit": "m", "note": "n"}},
        "names": {"n1": {"text": "hello"}},
        "tags": {"t1": {"name": "t1", "note": ""}},
        "tiles": {"x": {"name": "X", "tags": "t1", "top": "#000", "side": "#111",
                        "footing": "walk", "cross": 10, "clear": 10, "note": ""}},
        "attributes": {a: {"name": a, "abbrev": a[:3].upper(), "note": ""}
                       for a in ("might", "agility", "endurance", "presence",
                                 "intellect", "willpower")},
        "skills": dict(
            (a[:3] + "_" + b[:3],
             {"name": a + "/" + b, "derives": a + ":1," + b + ":1", "note": ""})
            for i, a in enumerate(("might", "agility", "endurance", "presence",
                                   "intellect", "willpower"))
            for b in ("might", "agility", "endurance", "presence", "intellect",
                      "willpower")[i + 1:]),
        "bones": {"root": {"name": "Root", "parent": "", "x": 0, "y": 0, "z": 0, "note": ""},
                  "chest": {"name": "Chest", "parent": "root", "x": 0, "y": 0, "z": 1, "note": ""}},
        "speeds": collections.OrderedDict([
            ("normal", {"name": "1x", "multiplier": 1, "note": ""}),
            ("fast", {"name": "2x", "multiplier": 2, "note": ""})]),
        "slots": dict((s_, {"name": s_, "note": ""}) for s_ in
                      ("head", "neck", "back", "torso", "gloves", "mainhand",
                       "offhand", "belt", "legs", "feet", "trinket1", "trinket2")),
        "gear": dict((s_, {"name": s_, "slot": s_, "tags": "t1", "note": ""}) for s_ in
                     ("head", "neck", "back", "torso", "gloves", "mainhand",
                      "offhand", "belt", "legs", "feet", "trinket1", "trinket2")),
        "figure": dict([("torso_body", {"name": "Torso", "slot": "body", "item": "",
                                        "bone": "chest", "colour": "#444",
                                        "ox": 0, "oy": 0, "from_m": -0.3, "to_m": 0.2,
                                        "w_top": 0.18, "w_bot": 0.14,
                                        "d_top": 0.1, "d_bot": 0.09,
                                        "sides": 6, "rings": 3, "bulge": 0.1,
                                        "cap_top": 0.1, "cap_bot": 0.1, "note": ""})]
                       + [(s_ + "_v", {"name": s_, "slot": s_, "item": s_,
                                       "bone": "chest", "colour": "#555",
                                       "ox": 0, "oy": 0, "from_m": -0.3, "to_m": 0.2,
                                       "w_top": 0.19, "w_bot": 0.15,
                                       "d_top": 0.11, "d_bot": 0.1,
                                       "sides": 6, "rings": 3, "bulge": 0.1,
                                       "cap_top": 0.1, "cap_bot": 0.1, "note": ""})
                          for s_ in ("head", "neck", "back", "torso", "gloves",
                                     "mainhand", "offhand", "belt", "legs",
                                     "feet", "trinket1", "trinket2")]),
        "structures": {"fire": {"name": "Fire", "tags": "t1", "colour": "#900",
                                "half_width": 0.3, "height_m": 0.5,
                                "skill": "mig_agi", "difficulty": 40, "note": ""}},
    }

    def sheet_from(base_, **patch):
        s = {k: {kk: dict(vv) for kk, vv in v.items()} for k, v in base_.items()}
        for path, val in patch.items():
            sec, key, col = path.split("__")
            s[sec][key][col] = val
        return s

    _, moved, errs = reconcile(base, sheet_from(base, knobs__a__value=5))
    add("a changed number is reported", moved == [("knobs", "a", "value", 1, 5)], str(moved))
    add("a changed number is not an error", not errs, str(errs))

    s = sheet_from(base); del s["knobs"]["a"]
    _, _, errs = reconcile(base, s)
    add("a knob missing from the sheet stops the build", len(errs) == 1, str(errs))

    s = sheet_from(base); s["knobs"]["ghost"] = {"value": 9, "unit": "", "note": ""}
    _, _, errs = reconcile(base, s)
    add("a dial the game never reads stops the build", len(errs) == 1, str(errs))

    _, _, errs = reconcile(base, sheet_from(base, geometry__g__value=99))
    add("edited geometry stops the build", len(errs) == 1, str(errs))

    _, moved, errs = reconcile(base, sheet_from(base, knobs__a__note="different prose"))
    add("an edited note is not a moved number", not moved and not errs, str(moved + errs))

    m, moved, errs = reconcile(base, sheet_from(base, knobs__a__value=None))
    add("an empty cell leaves the default alone",
        not moved and m["knobs"]["a"]["value"] == 1, str(m["knobs"]["a"]))

    m, _, _ = reconcile(base, sheet_from(base, tiles__x__tags="t1,nope"))
    add("a tile claiming an unknown tag stops the build",
        len(check_vocabulary(for_game(m))) == 1, "")

    m, _, _ = reconcile(base, sheet_from(base, skills__mig_agi__derives="agility:1,charm:1"))
    errs = check_vocabulary(for_game(m))
    add("a skill from an attribute nobody has stops the build (rule 2)",
        any("not one of the six attributes" in e for e in errs), str(errs[:1]))

    s4 = sheet_from(base)
    s4["figure"] = {k: v for k, v in s4["figure"].items() if k != "head_v"}
    d4 = {k: dict(v) for k, v in base.items()}
    d4["figure"] = {k: v for k, v in base["figure"].items() if k != "head_v"}
    m, _, _ = reconcile(d4, s4)
    add("gear nothing draws stops the build (rule 4)",
        any("would show nothing" in e for e in check_vocabulary(for_game(m))),
        str(check_vocabulary(for_game(m))[:1]))

    s5 = sheet_from(base, speeds__normal__multiplier=3)
    m, _, _ = reconcile(base, s5)
    add("normal speed must really be normal",
        any("must be 1" in e for e in check_vocabulary(for_game(m))),
        str(check_vocabulary(for_game(m))[:1]))

    m, _, _ = reconcile(base, None)
    add("twelve gear slots, no more and no less",
        not any("twelve" in e for e in check_vocabulary(for_game(m))), "")

    m, _, _ = reconcile(base, None)
    add("six attributes, and all fifteen pairs present (rule 2)",
        not check_vocabulary(for_game(m)), str(check_vocabulary(for_game(m))))

    s2 = sheet_from(base)
    del s2["skills"]["mig_agi"]
    d2 = {k: dict(v) for k, v in base.items()}
    d2["skills"] = {k: v for k, v in base["skills"].items() if k != "mig_agi"}
    m, _, _ = reconcile(d2, s2)
    add("a missing attribute pair stops the build",
        any("might + agility" in e or "agility + might" in e
            for e in check_vocabulary(for_game(m))),
        str(check_vocabulary(for_game(m))[:1]))

    m, _, _ = reconcile(base, None)
    add("with no sheet at all, the defaults still build",
        for_game(m)["knobs"]["a"] == 1, "")

    return checks
