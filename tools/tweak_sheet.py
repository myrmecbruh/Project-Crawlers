#!/usr/bin/env python3
"""Write a copy of the master sheet with one value changed.

    python3 tools/tweak_sheet.py <out.xlsx> <section> <key> <column> <value>

Only used by the tests, to prove end to end that turning a dial in the
spreadsheet really does change what appears on screen. It never touches
docs/crawlers.xlsx.
"""
import pathlib
import shutil
import sys

ROOT = pathlib.Path(__file__).resolve().parent.parent
SRC = ROOT / "docs" / "crawlers.xlsx"


def main():
    out, section, key, column, value = sys.argv[1:6]
    shutil.copy(SRC, out)
    from openpyxl import load_workbook
    wb = load_workbook(out)
    ws = wb[section]
    head = [str(c.value).strip().lower() if c.value else "" for c in ws[1]]
    col = head.index(column) + 1
    for row in range(2, ws.max_row + 1):
        if str(ws.cell(row=row, column=1).value).strip() == key:
            try:
                ws.cell(row=row, column=col, value=float(value)
                        if "." in value else int(value))
            except ValueError:
                ws.cell(row=row, column=col, value=value)
            wb.save(out)
            print("set %s.%s.%s = %s" % (section, key, column, value))
            return 0
    print("no such key: %s.%s" % (section, key), file=sys.stderr)
    return 1


if __name__ == "__main__":
    sys.exit(main())
